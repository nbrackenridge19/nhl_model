"""
nhl_migrate_arenas.py — one-off migration of arena capacity data.

The 'Arenas' sheet in nhl2526.xlsx already covers all 9 seasons (2017-18
through 2025-26) in one table, so this reads once from that single file
rather than looping through every season's own workbook.

Feeds HomeAdjD in v_team_game_features: HomeAdj collapses to just the
home team's (running-average-actual-attendance / capacity) ratio, since
the away team's Home flag is always 0.
"""

import os
import openpyxl
import psycopg2
from psycopg2.extras import execute_values

SOURCE_FILE = os.environ.get("NHL_SOURCE_FILE", "nhl2526.xlsx")
DRY_RUN = os.environ.get("DRY_RUN", "true").strip().lower() != "false"


def get_db_conn():
    if os.environ.get("PGHOST"):
        return psycopg2.connect()
    raise RuntimeError("Set PGHOST/PGUSER/PGPASSWORD/PGDATABASE/PGPORT.")


def main():
    wb = openpyxl.load_workbook(SOURCE_FILE, data_only=True)
    ws = wb["Arenas"]
    by_key = {}
    for r in range(2, ws.max_row + 1):
        season, team, cap = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value, ws.cell(row=r, column=6).value
        if not (season and team and cap):
            continue
        key = (str(season), team)
        # Some team/season pairs have a second row for a one-off outdoor
        # game (Winter Classic/Stadium Series) at a much larger stadium —
        # confirmed against real arena capacities (e.g. Buffalo's actual
        # home arena is 19,070; a second ~41,821 row is an outdoor game).
        # Keep the smaller (regular home arena) value.
        if key not in by_key or int(cap) < by_key[key]:
            by_key[key] = int(cap)

    rows = [(season, team, cap) for (season, team), cap in by_key.items()]
    print(f"Extracted {len(rows)} arena rows from {SOURCE_FILE} (regular-arena capacity only, "
          f"outdoor-game venues excluded)")

    if DRY_RUN:
        print("[dry-run] would upsert", len(rows), "rows")
        return

    conn = get_db_conn()
    try:
        sql = ("INSERT INTO arenas (season, team_code, cap) VALUES %s "
               "ON CONFLICT (season, team_code) DO UPDATE SET cap = excluded.cap")
        with conn.cursor() as cur:
            execute_values(cur, sql, rows)
        conn.commit()
        print(f"[live] arenas: upserted {len(rows)} rows")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
